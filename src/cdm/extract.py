"""Clinical text to OMOP fields with an LLM, scored against hand labels (research Phase G5).

From a note, three fields: the lesion's **diagnosis**, its **body site**, and whether it was
**confirmed by histopathology**. Each answer is a source label from the Phase A concept map
(``config/omop_concepts.csv``), so it lands on an OMOP concept ID, or "other" (stated, but no
mapped label fits) or "not stated". Two extractors answer the same question:

- :func:`llm_extract`: Claude Haiku through the Messages API (this model takes no temperature
  setting, so answers are cached and replayed rather than regenerated), with the
  allowed labels in the prompt. An answer outside them is kept as "invalid" and scored wrong.
- :func:`keyword_extract`: a keyword baseline, the bar the LLM has to clear.

Notes come from a :class:`NoteSource`: MTSamples now, MIMIC-IV-Note later, behind the same
interface. **Note text never enters the repository.** The labelling sheet, the LLM's raw
answers and the error sheet are written to a folder outside it (:func:`outside_repo`), and the
committed results hold aggregates only. Every LLM call is logged with its token cost, and a
budget guard stops calls before the cap.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import urllib.error
import urllib.request
from collections import Counter
from collections.abc import Callable, Iterable, Iterator
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Protocol

import numpy as np
import pandas as pd

REPO = Path(__file__).resolve().parents[2]
FIELDS = ("diagnosis", "body_site", "confirmation")
OTHER, NOT_STATED, INVALID = "other", "not stated", "invalid"
HISTOPATHOLOGY = "histopathology"
MODEL = "claude-haiku-5-5"
# US dollars per million tokens, prompts up to 100k tokens (Anthropic pricing page, read
# 2026-10-21). Used for the spend log and the budget guard.
PRICE_PER_MTOK = {MODEL: (0.10, 0.50)}
MAX_OUTPUT_TOKENS = 200


# ---- notes ----------------------------------------------------------------------------------


@dataclass(frozen=True)
class Note:
    note_id: str
    source: str
    specialty: str
    title: str
    text: str


def note_id(text: str) -> str:
    """A stable ID from the text, so labels and predictions join without storing text."""
    return hashlib.sha256(" ".join(text.split()).encode()).hexdigest()[:12]


class NoteSource(Protocol):
    """Anything that yields notes. Swap MTSamples for MIMIC-IV-Note by passing another one."""

    name: str

    def notes(self) -> Iterator[Note]: ...


SKIN = re.compile(
    r"\b(skin|cutaneous|derm\w*|melan\w*|nev(us|i)|basal cell|squamous cell|keratos\w*|"
    r"lentigo|epiderm\w*|mole)\b",
    re.I,
)
# Owner's decision, 2026-10-22, before any labelling (results/research/g5_extraction/
# selection_spec.json). SPECIFIC names a skin-lesion diagnosis or a skin biopsy; GENERIC only
# says "skin lesion" (often "no skin lesions" in a review of systems). Unlike SKIN, neither
# matches incidental words ("the skin was closed").
SPECIFIC = re.compile(
    r"\b(melanomas?|basal cell (carcinomas?|cancers?)|bcca|"
    r"squamous cell (carcinomas?|cancers?) of the skin|cutaneous squamous cell|keratoacanthomas?|"
    r"nev(us|i)|seborrheic keratos[ie]s|actinic keratos[ie]s|dermatofibromas?|cherry angiomas?|"
    r"pyogenic granulomas?|lentigo|lentigines|skin cancers?|pigmented lesions?|"
    r"skin biops(y|ies)|biops(y|ies) of the skin)\b",
    re.I,
)
GENERIC = re.compile(r"\b(skin|cutaneous) lesions?\b", re.I)
DEV_SHARE = 0.1
STRATA = ("dermatology", "specific_mention", "generic_mention_only")


def stratum(note: Note) -> str:
    """Which selection group a note is in; the development split takes from each."""
    if note.specialty == "Dermatology":
        return "dermatology"
    return "specific_mention" if SPECIFIC.search(note.text) else "generic_mention_only"


@dataclass
class MTSamples:
    """MTSamples transcriptions (mtsamples.com; the common CSV has ``medical_specialty``,
    ``sample_name`` and ``transcription``). Dermatology notes, pathology notes that are
    about skin, and notes from any other specialty that name a skin-lesion diagnosis, a skin
    biopsy (``SPECIFIC``) or a skin lesion (``GENERIC``). Duplicate texts count once, under
    the Dermatology row where one exists. The site allows sharing for education with credit,
    not redistribution under an open license, so the text stays outside the repository."""

    path: Path
    name: str = "mtsamples"

    def notes(self) -> Iterator[Note]:
        frame = pd.read_csv(self.path, dtype=str, keep_default_na=False)
        for _, row in frame.iterrows():
            specialty, text = row["medical_specialty"].strip(), row["transcription"].strip()
            if not text:
                continue
            derm = specialty == "Dermatology"
            skin_pathology = "pathology" in specialty.lower() and SKIN.search(text)
            if derm or skin_pathology or SPECIFIC.search(text) or GENERIC.search(text):
                yield Note(note_id(text), self.name, specialty, row["sample_name"].strip(), text)


@dataclass
class MimicIVNotes:
    """MIMIC-IV-Note (PhysioNet, credentialed access): ``discharge.csv.gz`` or
    ``radiology.csv.gz`` with ``note_id`` and ``text``. Tier 2 data
    (``docs/governance.md``): it is read only where the data use agreement allows, and only
    aggregates leave that environment. Notes that mention skin lesions are kept."""

    path: Path
    name: str = "mimic-iv-note"

    def notes(self) -> Iterator[Note]:
        if os.environ.get("CDM_TIER2") != "1":
            raise PermissionError("MIMIC-IV-Note is tier 2: set CDM_TIER2=1 where the DUA allows")
        outside_repo(self.path)
        for chunk in pd.read_csv(self.path, dtype=str, keep_default_na=False, chunksize=10_000):
            for _, row in chunk.iterrows():
                if SKIN.search(row["text"]):
                    yield Note(row["note_id"], self.name, "", "", row["text"])


def outside_repo(path: Path) -> Path:
    """Refuse a path inside this repository or any other git checkout (a second worktree, a
    clone): note text, raw answers and error sheets live elsewhere."""
    resolved = path.resolve()
    for folder in (resolved, *resolved.parents):
        if folder == REPO or (folder / ".git").exists():
            raise ValueError(f"{path} is inside the repository; note text must stay outside it")
    return path


# ---- the labels an answer may use -----------------------------------------------------------


@dataclass(frozen=True)
class Choices:
    """Allowed labels per field, each with its OMOP concept ID."""

    concept: dict[str, dict[str, int]]

    def labels(self, field: str) -> list[str]:
        return [*self.concept[field], OTHER, NOT_STATED]

    def value(self, field: str, answer: str) -> str:
        """What is scored: the concept ID for a mapped label, else "other", "not stated" or
        "invalid". Labels with the same concept score the same."""
        answer = answer.strip()
        if answer in (OTHER, NOT_STATED):
            return answer
        found = self.concept[field].get(answer)
        return str(found) if found is not None else INVALID


def choices(cmap: pd.DataFrame) -> Choices:
    """From the concept map: ISIC diagnosis labels and body sites (both levels) that have a
    concept, and histopathology. Confirmation methods without a concept (consensus, serial
    imaging) are concept 0 in the tables, so for scoring they are "not stated"."""

    def pick(fields: tuple[str, ...]) -> dict[str, int]:
        rows = cmap[cmap["source_field"].isin(fields) & cmap["scope"].isin(["all", "sites"])]
        rows = rows[rows["concept_id"].astype(int) != 0]
        return {r["source_value"]: int(r["concept_id"]) for _, r in rows.iterrows()}

    histo = pick(("diagnosis_confirm_type",))
    return Choices(
        {
            "diagnosis": pick(("diagnosis_3",)),
            "body_site": pick(("anatom_site_1", "anatom_site_2")),
            "confirmation": {HISTOPATHOLOGY: histo[HISTOPATHOLOGY]},
        }
    )


# ---- the LLM extractor ----------------------------------------------------------------------

SYSTEM = """You read one clinical note and fill three fields about the skin lesion the note \
is mainly about. Answer only with labels from the lists given, copied exactly.

- diagnosis: the lesion's final diagnosis. If the note states a diagnosis that is not in the \
list, answer "other". If it states none, answer "not stated". Do not guess from symptoms.
- body_site: where the lesion is. Prefer the most specific label that the note supports. \
"other" if the site is stated but not in the list; "not stated" if none.
- confirmation: "histopathology" if the diagnosis rests on a biopsy, excision or other \
specimen examined under the microscope; otherwise "not stated".

Reply with one JSON object and nothing else: \
{"diagnosis": "...", "body_site": "...", "confirmation": "..."}"""


def prompt(note: Note, allowed: Choices) -> str:
    lists = "\n".join(f"{f} labels: {json.dumps(allowed.labels(f))}" for f in FIELDS)
    return f"{lists}\n\nNote:\n<note>\n{note.text}\n</note>"


def prompt_hash(allowed: Choices) -> str:
    """Identifies the prompt version: the instructions and the allowed labels."""
    blob = SYSTEM + json.dumps({f: allowed.labels(f) for f in FIELDS})
    return hashlib.sha256(blob.encode()).hexdigest()[:12]


class BudgetExceeded(RuntimeError):
    """The next call could take spend past the cap, so it was not made."""


@dataclass
class SpendLog:
    """One JSON line per LLM call, with tokens and cost. Lives outside the repository (it
    names notes by ID only, but it is a working log, not a result)."""

    path: Path
    cap_usd: float = 5.0

    def entries(self) -> list[dict[str, Any]]:
        if not self.path.exists():
            return []
        return [json.loads(x) for x in self.path.read_text("utf-8").splitlines() if x.strip()]

    def total(self) -> float:
        return float(sum(e["cost_usd"] for e in self.entries()))

    def check(self, model: str, prompt_chars: int) -> None:
        """Refuse a call whose worst case (about 3 characters a token, the full output) would
        pass the cap."""
        price_in, price_out = PRICE_PER_MTOK[model]
        worst = (prompt_chars / 3 * price_in + MAX_OUTPUT_TOKENS * price_out) / 1e6
        if self.total() + worst > self.cap_usd:
            raise BudgetExceeded(f"spent ${self.total():.4f} of ${self.cap_usd:.2f}")

    def record_failure(self, model: str, note: str, prompt_sha: str, prompt_chars: int) -> None:
        """A call that failed after it may have been billed (a timeout, a dropped connection):
        logged at its worst-case cost, so the cap stays conservative."""
        price_in, price_out = PRICE_PER_MTOK[model]
        worst = (prompt_chars / 3 * price_in + MAX_OUTPUT_TOKENS * price_out) / 1e6
        self.path.parent.mkdir(parents=True, exist_ok=True)
        entry = {"ts": datetime.now(UTC).isoformat(timespec="seconds"), "model": model,
                 "note_id": note, "prompt_sha": prompt_sha, "failed": True,
                 "cost_usd": worst}  # fmt: skip
        with self.path.open("ab") as f:
            f.write((json.dumps(entry) + "\n").encode())

    def record(self, model: str, note: str, prompt_sha: str, usage: dict[str, int]) -> float:
        price_in, price_out = PRICE_PER_MTOK[model]
        cost = (usage["input_tokens"] * price_in + usage["output_tokens"] * price_out) / 1e6
        self.path.parent.mkdir(parents=True, exist_ok=True)
        entry = {"ts": datetime.now(UTC).isoformat(timespec="seconds"), "model": model,
                 "note_id": note, "prompt_sha": prompt_sha, **usage, "cost_usd": cost}  # fmt: skip
        with self.path.open("ab") as f:
            f.write((json.dumps(entry) + "\n").encode())
        return cost


# A transport takes (system, user message) and returns (reply text, usage). Tests pass a
# recorded one; :func:`anthropic_transport` calls the API.
Transport = Callable[[str, str], tuple[str, dict[str, int]]]


def api_key() -> str:
    """ANTHROPIC_API_KEY from the environment, or from the repository's .env (gitignored)."""
    key = os.environ.get("ANTHROPIC_API_KEY", "")
    env = REPO / ".env"
    if not key and env.exists():
        for line in env.read_text("utf-8").splitlines():
            if line.startswith("ANTHROPIC_API_KEY="):
                key = line.split("=", 1)[1].strip().strip('"').strip("'")
    if not key:
        raise RuntimeError("ANTHROPIC_API_KEY is not set (environment or .env)")
    return key


def anthropic_transport(model: str = MODEL) -> Transport:
    def send(system: str, user: str) -> tuple[str, dict[str, int]]:
        body = {"model": model, "max_tokens": MAX_OUTPUT_TOKENS, "system": system,
                "messages": [{"role": "user", "content": user}]}  # fmt: skip
        request = urllib.request.Request(
            "https://api.anthropic.com/v1/messages",
            data=json.dumps(body).encode(),
            headers={
                "x-api-key": api_key(),
                "anthropic-version": "2023-06-01",
                "content-type": "application/json",
            },  # fmt: skip
        )
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                reply = json.loads(response.read())
        except urllib.error.HTTPError as exc:  # the API's message says what was wrong
            raise RuntimeError(f"Messages API {exc.code}: {exc.read().decode()[:300]}") from exc
        text = "".join(part.get("text", "") for part in reply["content"])
        usage = {k: int(reply["usage"][k]) for k in ("input_tokens", "output_tokens")}
        return text, usage

    return send


def parse_reply(text: str) -> dict[str, str]:
    """The three answers from a reply; a field that is missing or not a string is
    "invalid"."""
    match = re.search(r"\{.*\}", text, re.S)
    try:
        raw = json.loads(match.group(0)) if match else {}
    except json.JSONDecodeError:
        raw = {}
    return {f: raw[f] if isinstance(raw.get(f), str) else INVALID for f in FIELDS}


def llm_extract(
    note: Note,
    allowed: Choices,
    send: Transport,
    spend: SpendLog,
    cache: Path,
    model: str = MODEL,
) -> dict[str, str]:
    """Haiku's answers for one note. A cached answer for the same note, model and prompt is
    reused, so a re-run costs nothing; the cache holds raw replies and stays outside the
    repository."""
    if note.source == MimicIVNotes.name:
        # PhysioNet's credentialed data use terms do not allow sending notes to an outside
        # model API; MIMIC-IV-Note runs only with a model inside the approved environment.
        raise PermissionError("MIMIC-IV-Note text is never sent to an outside LLM API")
    sha = prompt_hash(allowed)
    key = f"{model}:{sha}:{note.note_id}"
    outside_repo(cache)
    cached = _read_cache(cache)
    if key not in cached:
        user = prompt(note, allowed)
        chars = len(SYSTEM) + len(user)
        spend.check(model, chars)
        try:
            text, usage = send(SYSTEM, user)
        except Exception:
            spend.record_failure(model, note.note_id, sha, chars)
            raise
        spend.record(model, note.note_id, sha, usage)
        cache.parent.mkdir(parents=True, exist_ok=True)
        with cache.open("ab") as f:
            f.write((json.dumps({"key": key, "reply": text}) + "\n").encode())
        cached[key] = text
    return parse_reply(cached[key])


def _read_cache(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    rows = [json.loads(x) for x in path.read_text("utf-8").splitlines() if x.strip()]
    return {r["key"]: r["reply"] for r in rows}


# ---- the keyword baseline -------------------------------------------------------------------

# Extra words for labels whose ISIC wording is rare in notes. Everything else matches its own
# label text (without ", NOS").
SYNONYMS: dict[str, list[str]] = {
    "Melanoma, NOS": ["melanoma", "malignant melanoma"],
    "Melanoma Invasive": ["invasive melanoma", "invasive malignant melanoma"],
    "Melanoma in situ": ["melanoma in situ", "lentigo maligna"],
    "Basal cell carcinoma": ["basal cell carcinoma", "basal cell epithelioma", "bcc"],
    "Squamous cell carcinoma, NOS": ["squamous cell carcinoma", "scc"],
    "Solar or actinic keratosis": ["actinic keratosis", "actinic keratoses", "solar keratosis"],
    "Nevus": ["nevus", "nevi", "naevus", "mole"],
    "Seborrheic keratosis": ["seborrheic keratosis", "seborrheic keratoses"],
    "Verruca": ["verruca", "wart", "warts"],
    "Hemangioma": ["hemangioma", "cherry angioma"],
    "Head and neck": ["head and neck"],
    "Lower extremity": ["lower extremity", "leg"],
    "Upper extremity": ["upper extremity", "arm"],
    "Posterior trunk": ["back", "posterior trunk"],
    "Anterior trunk": ["chest", "abdomen", "anterior trunk"],
}
HISTO_WORDS = re.compile(
    r"\b(biops\w*|patholog\w*|histolog\w*|histopatholog\w*|microscopic|frozen section|"
    r"specimen|excis\w*)\b",
    re.I,
)


def _patterns(label: str) -> list[re.Pattern[str]]:
    words = SYNONYMS.get(label, [label.replace(", NOS", "")])
    return [re.compile(r"\b" + re.escape(w) + r"\b", re.I) for w in words]


def keyword_extract(note: Note, allowed: Choices) -> dict[str, str]:
    """First mention wins; at the same position, the longer (more specific) label. No mapped
    label found is "not stated": keywords cannot tell "other" apart."""
    out = {}
    for field in ("diagnosis", "body_site"):
        hits = []
        for label in allowed.concept[field]:
            for pattern in _patterns(label):
                found = pattern.search(note.text)
                if found:
                    hits.append((found.start(), -len(found.group(0)), label))
        out[field] = min(hits)[2] if hits else NOT_STATED
    out["confirmation"] = HISTOPATHOLOGY if HISTO_WORDS.search(note.text) else NOT_STATED
    return out


# ---- the labelling sheet --------------------------------------------------------------------

SHEET_COLUMNS = ["note_id", "split", "source", "specialty", "title", "text",
                 *(f"label_{f}" for f in FIELDS), "labeller_notes"]  # fmt: skip


def split_notes(
    notes: Iterable[Note], n: int = 100, dev_share: float = DEV_SHARE
) -> list[tuple[Note, str]]:
    """Up to ``n`` notes in a fixed order (by note ID). Within each selection group
    (:func:`stratum`), the first ``dev_share`` of the group (at least one) are for prompt design
    and the rest for scoring, so the development notes are not all of one kind. Decided before
    any prompt runs, as the protocol requires (docs/EVAL_PLAN.md, section 3)."""
    unique: dict[str, Note] = {}
    for x in notes:
        # A text listed under several specialties counts once, under Dermatology if it is there.
        if x.note_id not in unique or x.specialty == "Dermatology":
            unique[x.note_id] = x
    chosen = sorted(unique.values(), key=lambda x: x.note_id)[:n]
    dev: set[str] = set()
    for group in STRATA:
        members = [x for x in chosen if stratum(x) == group]
        dev |= {x.note_id for x in members[: max(1, int(len(members) * dev_share + 0.5))]}
    return [(x, "dev" if x.note_id in dev else "scoring") for x in chosen]


def selection_summary(notes: Iterable[Note], n: int = 100) -> dict[str, Any]:
    """Aggregates of the selection, for the committed spec: counts only, never note text. The
    hash is of the sorted note IDs joined by newlines."""
    split = split_notes(notes, n)
    by = {g: {"dev": 0, "scoring": 0} for g in STRATA}
    for note, s in split:
        by[stratum(note)][s] += 1
    ids = "\n".join(sorted(note.note_id for note, _ in split))
    return {
        "notes": len(split),
        "by_split": {s: sum(1 for _, x in split if x == s) for s in ("dev", "scoring")},
        "by_group_and_split": by,
        "note_ids_sha256": hashlib.sha256(ids.encode()).hexdigest(),
    }


def write_sheet(notes: Iterable[Note], folder: Path, allowed: Choices, n: int = 100) -> Path:
    """The labelling sheet (CSV, with the note text: outside the repository) and a second
    file listing the labels allowed in each column. An existing sheet is never overwritten:
    it may hold your labels, and a new one could move notes between dev and scoring."""
    outside_repo(folder).mkdir(parents=True, exist_ok=True)
    sheet = folder / "labelling_sheet.csv"
    if sheet.exists():
        raise FileExistsError(f"{sheet} exists; move it away first if you mean to start over")
    with sheet.open("w", newline="", encoding="utf-8") as f:
        rows = csv.DictWriter(f, SHEET_COLUMNS)
        rows.writeheader()
        for note, split in split_notes(notes, n):
            blank = {f"label_{f}": "" for f in FIELDS} | {"labeller_notes": ""}
            rows.writerow({"note_id": note.note_id, "split": split, "source": note.source,
                           "specialty": note.specialty, "title": note.title, "text": note.text,
                           **blank})  # fmt: skip
    with (folder / "labelling_choices.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["column", "allowed label", "OMOP concept ID"])
        for field in FIELDS:
            for label in allowed.labels(field):
                writer.writerow([f"label_{field}", label, allowed.concept[field].get(label, "")])
    return sheet


# ---- scoring --------------------------------------------------------------------------------

# Labels whose concepts are one family: a confusion inside a family is "wrong granularity".
FAMILIES = [
    {"Melanoma, NOS", "Melanoma Invasive", "Melanoma in situ"},
    {"Squamous cell carcinoma, NOS", "Squamous cell carcinoma in situ"},
    {"Head and neck", "Head", "Neck"},
    {"Trunk", "Anterior trunk", "Posterior trunk", "Lateral trunk", "Buttock"},
    {"Upper extremity", "Upper arm", "Elbow", "Forearm", "Wrist", "Hand", "Shoulder"},
    {"Lower extremity", "Thigh", "Knee", "Lower leg", "Ankle", "Foot"},
]


def error_kind(field: str, label: str, predicted: str, allowed: Choices) -> str | None:
    """None if correct, else: missed, hallucinated, wrong granularity, wrong concept, or
    invalid. "Label error" is the labeller's call, made on the error sheet."""
    truth, guess = allowed.value(field, label), allowed.value(field, predicted)
    if guess == INVALID:  # never correct, even against a label that is itself invalid
        return "invalid"
    if truth == guess:
        return None
    if guess == NOT_STATED:
        return "missed"
    if truth == NOT_STATED:
        return "hallucinated"
    if any(label in fam and predicted in fam for fam in FAMILIES):
        return "wrong granularity"
    return "wrong concept"


def _interval(hits: np.ndarray, n: int = 2000, seed: int = 0) -> dict[str, float]:
    rng = np.random.default_rng(seed)
    means = [hits[rng.integers(0, len(hits), len(hits))].mean() for _ in range(n)]
    return {"low": float(np.percentile(means, 2.5)), "high": float(np.percentile(means, 97.5))}


def score(
    labels: pd.DataFrame, predictions: dict[str, dict[str, str]], allowed: Choices
) -> dict[str, Any]:
    """Per field: accuracy on OMOP concept (with a 95% interval from resampling notes),
    precision and recall of "stated", and error counts by kind. Aggregates only."""
    out: dict[str, Any] = {"notes": len(labels)}
    for field in FIELDS:
        rows = [(r[f"label_{field}"], predictions[r["note_id"]][field])
                for _, r in labels.iterrows()]  # fmt: skip
        hits = np.array([error_kind(field, t, p, allowed) is None for t, p in rows], dtype=float)
        stated_true = np.array([allowed.value(field, t) != NOT_STATED for t, _ in rows])
        stated_pred = np.array([allowed.value(field, p) != NOT_STATED for _, p in rows])
        both = int((stated_true & stated_pred).sum())
        out[field] = {
            "accuracy": float(hits.mean()) if len(hits) else None,
            "interval": _interval(hits) if len(hits) else None,
            "stated_precision": both / int(stated_pred.sum()) if stated_pred.any() else None,
            "stated_recall": both / int(stated_true.sum()) if stated_true.any() else None,
            "labelled_stated": int(stated_true.sum()),
            "errors": dict(sorted(Counter(
                k for t, p in rows if (k := error_kind(field, t, p, allowed))).items())),
        }  # fmt: skip
    return out


def write_workbook(folder: Path) -> Path:
    """The labelling sheet as an Excel workbook beside the CSV (outside the repository): one
    dropdown per label column from ``labelling_choices.csv``, wrapped note text, and a protected
    sheet in which only the label and notes columns can be edited. The CSV stays the file the
    scorer reads; copy the labels back into it after labelling."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    sheet = outside_repo(folder) / "labelling_sheet.csv"
    rows = list(csv.DictReader(sheet.open(encoding="utf-8", newline="")))
    choices: dict[str, list[str]] = {}
    for row in csv.DictReader(
        (folder / "labelling_choices.csv").open(encoding="utf-8", newline="")
    ):
        choices.setdefault(row["column"], []).append(row["allowed label"])
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    lists = wb.create_sheet("choices")
    lists.sheet_state = "hidden"
    ws.append(SHEET_COLUMNS)
    for row in rows:
        ws.append([row[c] for c in SHEET_COLUMNS])
    editable = [c for c in SHEET_COLUMNS if c.startswith("label_") or c == "labeller_notes"]
    widths = {"text": 90, "title": 28, "labeller_notes": 30}
    for j, name in enumerate(SHEET_COLUMNS, start=1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = widths.get(name, 22 if name in editable else 14)
        for i in range(2, len(rows) + 2):
            cell = ws.cell(row=i, column=j)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.protection = Protection(locked=name not in editable)
        if name in choices:
            col = get_column_letter(list(choices).index(name) + 1)
            for r, value in enumerate(choices[name], start=1):
                lists[f"{col}{r}"] = value
            span = f"=choices!${col}$1:${col}${len(choices[name])}"
            check = DataValidation(
                type="list", formula1=span, allow_blank=True, showErrorMessage=True
            )
            check.error = "Pick a value from the list (see labelling_choices.csv)."
            ws.add_data_validation(check)
            check.add(f"{letter}2:{letter}{len(rows) + 1}")
    ws.freeze_panes = "A2"
    ws.protection.sheet = True  # the note text and IDs stay as written
    out = folder / "labelling_sheet.xlsx"
    if out.exists():
        raise FileExistsError(f"{out} exists; it may hold labels")
    wb.save(out)
    return out
